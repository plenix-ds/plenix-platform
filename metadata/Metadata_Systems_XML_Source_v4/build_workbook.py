"""
Build the master Excel workbook from XML files.
- One tab per system
- One row per attribute (or enumerated value, child rows)
- Index tab with summary
- Pending systems get a placeholder tab
"""
import xml.etree.ElementTree as ET
from pathlib import Path
import sys
sys.path.insert(0, "build")
from inventory import SYSTEMS, CATEGORIES, CAT_COLORS, sheet_name

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

XML_DIR = Path("xml")
OUT = Path("Metadata_Systems_Inventory_v2.xlsx")

# Styles
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", start_color="1F3864")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
BODY_FONT = Font(name="Calibri", size=10)
BODY_FONT_BOLD = Font(name="Calibri", size=10, bold=True)
WRAP_TOP = Alignment(vertical="top", wrap_text=True)
CENTER_TOP = Alignment(horizontal="center", vertical="top", wrap_text=True)
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

ENUM_FILL = PatternFill("solid", start_color="F7F7F7")
ENUM_FONT = Font(name="Calibri", size=9, color="404040", italic=True)


def style_header_row(ws, row_idx, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER


def write_system_sheet(wb, sid, name, used_for, cat, xml_path):
    """Write one tab for a system. xml_path may be None for pending systems."""
    tab = sheet_name(sid, name)
    ws = wb.create_sheet(tab)

    cat_label = CATEGORIES[cat]
    cat_color = CAT_COLORS[cat]

    # Top header block (rows 1-5): system metadata
    ws.cell(row=1, column=1, value=f"#{sid:03d}  {name}")
    ws.cell(row=1, column=1).font = Font(name="Calibri", size=14, bold=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)

    ws.cell(row=2, column=1, value="Category:")
    ws.cell(row=2, column=1).font = BODY_FONT_BOLD
    ws.cell(row=2, column=2, value=f"{cat} — {cat_label}")
    ws.cell(row=2, column=2).fill = PatternFill("solid", start_color=cat_color)

    ws.cell(row=3, column=1, value="Used for:")
    ws.cell(row=3, column=1).font = BODY_FONT_BOLD
    ws.cell(row=3, column=2, value=used_for)
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=6)

    if xml_path is None:
        # Pending placeholder
        ws.cell(row=4, column=1, value="Status:")
        ws.cell(row=4, column=1).font = BODY_FONT_BOLD
        ws.cell(row=4, column=2, value="PENDING — to be authored in a subsequent session")
        ws.cell(row=4, column=2).font = Font(name="Calibri", size=10, italic=True, color="C00000")
        ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=6)
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 60
        ws.column_dimensions["C"].width = 30
        return 0, 0

    tree = ET.parse(xml_path)
    root = tree.getroot()
    spec_source = root.get("spec_source", "")
    completeness = root.get("completeness", "")
    record_count = root.get("record_count", "")

    ws.cell(row=4, column=1, value="Spec source:")
    ws.cell(row=4, column=1).font = BODY_FONT_BOLD
    ws.cell(row=4, column=2, value=spec_source)
    ws.cell(row=4, column=2).font = Font(name="Calibri", size=9, italic=True)
    ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=6)

    ws.cell(row=5, column=1, value="Completeness:")
    ws.cell(row=5, column=1).font = BODY_FONT_BOLD
    ws.cell(row=5, column=2, value=f"{completeness} ({record_count} attributes)")
    ws.merge_cells(start_row=5, start_column=2, end_row=5, end_column=6)

    desc_el = root.find("description")
    if desc_el is not None and desc_el.text:
        ws.cell(row=6, column=1, value="Description:")
        ws.cell(row=6, column=1).font = BODY_FONT_BOLD
        ws.cell(row=6, column=2, value=desc_el.text.strip())
        ws.cell(row=6, column=2).alignment = WRAP_TOP
        ws.merge_cells(start_row=6, start_column=2, end_row=6, end_column=6)
        ws.row_dimensions[6].height = 45
        header_row = 8
    else:
        header_row = 7

    # Attribute table headers
    headers = ["#", "Name", "Tag/ID", "Type", "Group", "Description"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=header_row, column=c, value=h)
    style_header_row(ws, header_row, len(headers))

    row = header_row + 1
    attrs = root.findall("attributes/attribute")
    enum_count = 0

    for i, attr in enumerate(attrs, 1):
        def get(t):
            el = attr.find(t)
            return (el.text or "").strip() if el is not None and el.text else ""

        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=get("name"))
        ws.cell(row=row, column=3, value=get("tag_id") or get("keyword"))
        ws.cell(row=row, column=4, value=get("data_type"))
        ws.cell(row=row, column=5, value=get("group") or get("subsystem"))
        # Description with optional supplementary fields
        desc_parts = [get("description")]
        for extra in ("required", "writable", "value_multiplicity", "status"):
            v = get(extra)
            if v:
                desc_parts.append(f"[{extra}: {v}]")
        ws.cell(row=row, column=6, value=" ".join(p for p in desc_parts if p))
        for c in range(1, 7):
            cell = ws.cell(row=row, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = CENTER_TOP if c == 1 else WRAP_TOP
        row += 1

        # Enum values as indented child rows
        values = attr.findall("values/value")
        for v in values:
            key = v.get("key", "")
            val = (v.text or "").strip()
            ws.cell(row=row, column=2, value=f"    = {key}")
            ws.cell(row=row, column=2).font = ENUM_FONT
            ws.cell(row=row, column=2).alignment = WRAP_TOP
            ws.cell(row=row, column=2).fill = ENUM_FILL
            ws.cell(row=row, column=6, value=val)
            ws.cell(row=row, column=6).font = ENUM_FONT
            ws.cell(row=row, column=6).alignment = WRAP_TOP
            ws.cell(row=row, column=6).fill = ENUM_FILL
            row += 1
            enum_count += 1

    # Column widths
    widths = {1: 6, 2: 30, 3: 16, 4: 20, 5: 18, 6: 70}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:F{row - 1}"

    return len(attrs), enum_count


def write_index(wb, stats):
    """Build the master index tab."""
    ws = wb.create_sheet("Index", 0)
    ws.cell(row=1, column=1, value="Metadata Systems Master Index")
    ws.cell(row=1, column=1).font = Font(name="Calibri", size=16, bold=True, color="1F3864")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)

    headers = ["#", "Category Code", "Category", "Metadata System", "Status", "Attributes", "Enum Values"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=3, column=c, value=h)
    style_header_row(ws, 3, len(headers))

    row = 4
    for sid, name, used_for, cat, *_ in SYSTEMS:
        s = stats.get(sid, {"status": "PENDING", "attrs": 0, "enums": 0})
        ws.cell(row=row, column=1, value=sid)
        ws.cell(row=row, column=2, value=cat)
        ws.cell(row=row, column=3, value=CATEGORIES[cat])
        ws.cell(row=row, column=4, value=name)
        ws.cell(row=row, column=5, value=s["status"])
        ws.cell(row=row, column=6, value=s["attrs"])
        ws.cell(row=row, column=7, value=s["enums"])

        for c in range(1, 8):
            cell = ws.cell(row=row, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = CENTER_TOP if c in (1, 2, 6, 7) else WRAP_TOP
            if c in (2, 3):
                cell.fill = PatternFill("solid", start_color=CAT_COLORS[cat])
        if s["status"] == "PENDING":
            ws.cell(row=row, column=5).font = Font(name="Calibri", size=10, italic=True, color="C00000")
        row += 1

    widths = {1: 6, 2: 14, 3: 30, 4: 42, 5: 16, 6: 12, 7: 12}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:G{row - 1}"


def write_summary(wb, stats):
    ws = wb.create_sheet("Category Summary", 1)
    ws.cell(row=1, column=1, value="Category Summary")
    ws.cell(row=1, column=1).font = Font(name="Calibri", size=16, bold=True, color="1F3864")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)

    headers = ["Code", "Category", "Systems", "Attributes", "Enum Values"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=3, column=c, value=h)
    style_header_row(ws, 3, len(headers))

    by_cat = {}
    for sid, name, used_for, cat, *_ in SYSTEMS:
        by_cat.setdefault(cat, {"systems": 0, "attrs": 0, "enums": 0})
        by_cat[cat]["systems"] += 1
        s = stats.get(sid, {"attrs": 0, "enums": 0})
        by_cat[cat]["attrs"] += s["attrs"]
        by_cat[cat]["enums"] += s["enums"]

    row = 4
    for cat in CATEGORIES:
        if cat not in by_cat:
            continue
        v = by_cat[cat]
        ws.cell(row=row, column=1, value=cat)
        ws.cell(row=row, column=2, value=CATEGORIES[cat])
        ws.cell(row=row, column=3, value=v["systems"])
        ws.cell(row=row, column=4, value=v["attrs"])
        ws.cell(row=row, column=5, value=v["enums"])
        for c in range(1, 6):
            cell = ws.cell(row=row, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = CENTER_TOP if c in (1, 3, 4, 5) else WRAP_TOP
            if c in (1, 2):
                cell.fill = PatternFill("solid", start_color=CAT_COLORS[cat])
        row += 1

    # Total
    ws.cell(row=row, column=2, value="TOTAL").font = BODY_FONT_BOLD
    ws.cell(row=row, column=3, value=sum(v["systems"] for v in by_cat.values())).font = BODY_FONT_BOLD
    ws.cell(row=row, column=4, value=sum(v["attrs"] for v in by_cat.values())).font = BODY_FONT_BOLD
    ws.cell(row=row, column=5, value=sum(v["enums"] for v in by_cat.values())).font = BODY_FONT_BOLD
    for c in range(1, 6):
        ws.cell(row=row, column=c).fill = PatternFill("solid", start_color="D9D9D9")
        ws.cell(row=row, column=c).border = BORDER
        ws.cell(row=row, column=c).alignment = CENTER_TOP if c != 2 else WRAP_TOP

    widths = {1: 10, 2: 32, 3: 12, 4: 14, 5: 14}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w


def main():
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Build map of XML files
    xml_files = {}
    for f in XML_DIR.glob("*.xml"):
        sid = int(f.stem.split("_")[0])
        xml_files[sid] = f

    stats = {}
    print(f"Building workbook with {len(SYSTEMS)} system tabs...")

    for sid, name, used_for, cat, *_ in SYSTEMS:
        xml = xml_files.get(sid)
        attrs, enums = write_system_sheet(wb, sid, name, used_for, cat, xml)
        if xml is not None:
            stats[sid] = {"status": "DONE", "attrs": attrs, "enums": enums}
        else:
            stats[sid] = {"status": "PENDING", "attrs": 0, "enums": 0}
        if sid % 30 == 0:
            print(f"  {sid}/240 tabs done")

    write_index(wb, stats)
    write_summary(wb, stats)

    # Reorder: Index, Summary, then 001-240
    ordering = ["Index", "Category Summary"] + [s for s in wb.sheetnames if s not in ("Index", "Category Summary")]
    wb._sheets = [wb[n] for n in ordering]

    print(f"Saving to {OUT}...")
    wb.save(OUT)
    print(f"Done. {OUT} ({OUT.stat().st_size / 1e6:.2f} MB)")


if __name__ == "__main__":
    main()
